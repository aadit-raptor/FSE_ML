"""A sector model as a live Excel workbook (PLAN.md R4).

The same rows the server works out, written as Excel formulas from the same
parsed trees: plain arithmetic, MAX, MIN and one IF (for `div`). Formulas use
ordinary cell addresses, so Excel's trace arrows work; every assumption cell
and constant also has a defined name (`s1_occ_y1`, `unit_rupees`) by which
the add-in will find it. Typed cells are blue; assumptions are yellow; a
marked figure is amber or red with its reason in a comment and on the Notes
sheet.

Nothing is worked out here: the workbook holds formulas without stored
values, and the spreadsheet program that opens it does the arithmetic.
"""
from __future__ import annotations

import re
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties

from research import formula
from research.definition import ASSUME, Company, Heading, Model, ModelError, cells, label, words_for
from research.evaluate import target
from research.formula import Ref

SOURCE = "variater.com"
SHEET = "Model"
FIRST_CONSTANT_ROW = 4
FIRST_PERIOD_COLUMN = 3          # A = the row's label, B = its unit

FORMATS = {
    "money": "#,##0.00;(#,##0.00);-", "count": "#,##0;(#,##0);-", "count3": "#,##0.000", "days": "0",
    "share": "0.0%;-0.0%;-", "rupees": "#,##0;(#,##0);-", "rupees2": "0.00", "times": "0.00", "check": "0.00;-0.00;0.00",
}
BLUE_RGB, GREY_RGB, WHITE_RGB, GREEN, RED = "0B4F9E", "888888", "FFFFFF", "1D6F42", "B03A2E"
HEAD_RGB, SECTION_RGB, ASSUMED_RGB = "1F2A30", "DCE6EA", "FFF8D6"
MARKED_RGB = {"amber": "F6C566", "red": "F2A097"}
_LOOKS_LIKE_A_CELL = re.compile(r"[A-Za-z]{1,3}[0-9]+|[Rr][0-9]*[Cc][0-9]*")
# A saved model's sheet says which company's model it is (PLAN.md R5), so "Save to Variater" knows where to save
SAVED_NAME, SAVED_AT = "variater_model", (1, 2)
SAVED_TEXT = "Saved in Variater: "


@dataclass(frozen=True)
class Layout:
    """Where every row, period and constant of a model stands on its sheet."""
    header_row: int
    rows: dict              # row key -> sheet row
    columns: dict           # period name -> column letter
    constants: dict         # constant name -> address

    def cell(self, row: str, period: str) -> str:
        return "%s%d" % (self.columns[period], self.rows[row])

    def constant(self, name: str) -> str:
        return self.constants[name]


def layout(model: Model) -> Layout:
    constants = {name: "C%d" % (FIRST_CONSTANT_ROW + i) for i, name in enumerate(model.constants)}
    header = FIRST_CONSTANT_ROW + len(model.constants) + 1
    rows = {r.key: header + 1 + i for i, r in enumerate(model.rows) if not isinstance(r, Heading)}
    columns = {p.name: get_column_letter(FIRST_PERIOD_COLUMN + i) for i, p in enumerate(model.periods)}
    return Layout(header, rows, columns, constants)


def _absolute(address: str) -> str:
    return re.sub(r"([A-Z]+)([0-9]+)", r"$\1$\2", address)


def _put(ws, row: int, column: int, text: str):
    """Words into a cell, as words: a name or label that starts with = must never become a formula."""
    cell = ws.cell(row, column)
    cell.value = text
    cell.data_type = "s"
    return cell


def _unit_text(unit: str, company: Company, words: dict) -> str:
    money = company.unit_name
    return {"money": money, "check": money, "share": "%", "rupees": words.get("currency", ""), "rupees2": words.get("currency", ""),
            "count": "number", "count3": "number", "days": "days", "times": "times"}[unit]


def sheets(model: Model, company: Company, allow_missing: bool = False, saved: str | None = None) -> list[dict]:
    """The model for one company as sheets of cells: the one description the
    download (`write`) and the Excel add-in's "Load model" both write, so the
    two can never hold different cells. The shape is research/coverage.py
    `layout`'s, plus a cell's `color`, `fill`, `size`, `align` and `wrap`,
    and a sheet's `names` (defined name -> address) and `heights`.

    With `allow_missing`, assumptions not typed yet are blank yellow cells:
    the analyst types them in Excel and the formulas do the rest. `saved`, a
    saved model's company id (R5), is written beside the title and named
    `variater_model`, so the add-in's "Save to Variater" finds the model.
    """
    plan = cells(model, company, allow_missing=allow_missing)
    words = words_for(model, company)
    where = layout(model)
    width = FIRST_PERIOD_COLUMN + len(model.periods) - 1
    out: list = []
    names: dict = {}

    def put(row, column, **spec):
        out.append({"row": row, "column": column, **{k: v for k, v in spec.items() if v is not None and v is not False}})

    def name(key: str, address: str) -> None:
        if _LOOKS_LIKE_A_CELL.fullmatch(key):
            raise ModelError("%r cannot be a name in Excel: it reads as a cell" % key)
        if key in names:
            raise ModelError("two cells would both be named %r" % key)
        names[key] = _absolute(address)

    title = "%s: %s model%s" % (company.name, model.name, " (%s)" % company.sector if company.sector else "")
    put(1, 1, text=title, bold=True, size=13)
    if saved is not None:
        put(*SAVED_AT, text=SAVED_TEXT + str(saved), color=GREY_RGB)
        name(SAVED_NAME, "B1")
    vetted = ("Read by a practising analyst." if model.vetted else
              "NOT YET VETTED: no practising analyst has read this model yet.")
    put(2, 1, text=vetted + " Yellow cells are assumptions to type; everything in black is a formula. Source: " + SOURCE,
        bold=True, color=GREEN if model.vetted else RED)
    put(3, 1, text="%sMoney in %s." % (company.basis + ". " if company.basis else "", company.unit_name))
    for key, text in model.constants.items():
        at = where.constant(key)
        put(int(at[1:]), 1, text=label(text, words))
        put(int(at[1:]), 3, number=float(company.constants[key]), color=BLUE_RGB, number_format="#,##0.######")
        name(key, at)

    head = dict(bold=True, color=WHITE_RGB, fill=HEAD_RGB, wrap=True)
    put(where.header_row, 1, text="Row", **head)
    put(where.header_row, 2, text="Unit", **head)
    for i, period in enumerate(model.periods):
        put(where.header_row, FIRST_PERIOD_COLUMN + i, text=str(company.period_labels.get(period.name, period.name)),
            align="right", **head)

    def address(ref: Ref, period: str) -> str:
        if ref.name in where.constants:
            return _absolute(where.constants[ref.name])
        row, column = target(model, ref, period)
        return where.cell(row, column)

    marked = []
    n = where.header_row
    for row in model.rows:
        n += 1
        text = label(row.label, words)
        if isinstance(row, Heading):
            put(n, 1, text=text, bold=True, fill=SECTION_RGB)
            for column in range(2, width + 1):
                put(n, column, fill=SECTION_RGB)
            continue
        put(n, 1, text=text, bold=row.bold)
        put(n, 2, text=_unit_text(row.unit, company, words), color=GREY_RGB)
        for i, period in enumerate(model.periods):
            held = plan.get((row.key, period.name))
            if held is None:
                continue
            spec = dict(number_format=FORMATS[row.unit], bold=row.bold)
            if held.tree is not None:
                spec["formula"] = "=" + formula.to_excel(held.tree, lambda ref, p=period.name: address(ref, p))
            else:
                spec.update(number=held.value, color=BLUE_RGB)
                if held.typed is ASSUME:
                    spec["fill"] = ASSUMED_RGB
                    name("%s_%s" % (row.key, period.name), where.cell(row.key, period.name))
                mark = company.marks.get(row.key, {}).get(period.name)
                if mark:
                    spec.update(fill=MARKED_RGB[mark["level"]], note=str(mark.get("why", "")))
                    marked.append(("%s!%s" % (SHEET, where.cell(row.key, period.name)),
                                   "%s, %s" % (text.strip(), company.period_labels.get(period.name, period.name)),
                                   mark["level"], str(mark.get("why", ""))))
            put(n, FIRST_PERIOD_COLUMN + i, **spec)

    widths = {"A": 74, "B": 12, **{where.columns[p.name]: 17 for p in model.periods}}
    model_sheet = {"title": SHEET, "cells": out, "widths": widths, "freeze": "C%d" % (where.header_row + 1),
                   "names": names, "heights": {str(where.header_row): 32}}
    return [model_sheet, _notes(model, marked)]


def _notes(model: Model, marked: list) -> dict:
    out: list = []

    def put(row, column, **spec):
        out.append({"row": row, "column": column, **{k: v for k, v in spec.items() if v is not None and v is not False}})

    put(1, 1, text="What this model does and leaves out", bold=True, size=13)
    n = 2
    for note in model.notes:
        n += 1
        put(n, 1, text=note)
    n += 2
    put(n, 1, text="Figures to look at again, and why", bold=True, size=13)
    n += 1
    for j, head in enumerate(["Cell", "Figure", "Level", "Why"]):
        put(n, 1 + j, text=head, bold=True, color=WHITE_RGB, fill=HEAD_RGB)
    for at, figure, level, why in marked:
        n += 1
        put(n, 1, text=at)
        put(n, 2, text=figure)
        put(n, 3, text=level, fill=MARKED_RGB[level])
        put(n, 4, text=why)
    if not marked:
        put(n + 1, 1, text="None.")
    return {"title": "Notes", "cells": out, "widths": {"A": 18, "B": 60, "C": 9, "D": 140}, "freeze": None,
            "names": {}, "heights": {}}


def write(model: Model, company: Company, allow_missing: bool = False, saved: str | None = None) -> Workbook:
    """The model for one company as a workbook, written from `sheets`: a Model sheet of live formulas and a Notes sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = wb.properties.lastModifiedBy = SOURCE
    wb.calculation = CalcProperties(fullCalcOnLoad=True)      # no values are stored: work everything out on opening
    for sheet in sheets(model, company, allow_missing=allow_missing, saved=saved):
        ws = wb.create_sheet(sheet["title"])
        for spec in sheet["cells"]:
            if "text" in spec:
                cell = _put(ws, spec["row"], spec["column"], spec["text"])
            else:
                cell = ws.cell(spec["row"], spec["column"])
                cell.value = spec.get("formula", spec.get("number"))
            if spec.get("number_format"):
                cell.number_format = spec["number_format"]
            if spec.get("bold") or spec.get("color") or spec.get("size"):
                cell.font = Font(bold=bool(spec.get("bold")), color=spec.get("color"), size=spec.get("size"))
            if spec.get("fill"):
                cell.fill = PatternFill("solid", fgColor=spec["fill"])
            if spec.get("align") or spec.get("wrap"):
                cell.alignment = Alignment(horizontal=spec.get("align"), wrap_text=bool(spec.get("wrap")))
            if spec.get("note"):
                cell.comment = Comment(spec["note"], SOURCE)
        for key, address in sheet["names"].items():
            wb.defined_names[key] = DefinedName(key, attr_text="%s!%s" % (sheet["title"], address))
        for letter, width in sheet["widths"].items():
            ws.column_dimensions[letter].width = width
        for row, height in sheet["heights"].items():
            ws.row_dimensions[int(row)].height = height
        if sheet["freeze"]:
            ws.freeze_panes = sheet["freeze"]
    return wb
