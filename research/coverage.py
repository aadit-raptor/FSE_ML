"""A read the analyst has reviewed, made into stored periods, and the checked
workbook (PLAN.md R2).

The results reader (``research/reader``) answers a job result: read A's
tables, every doubtful figure amber with its reason. **Nothing is stored
until the analyst accepts it.** ``accept`` takes that result and the
analyst's review -- each amber figure either accepted as read or corrected,
and any other figure corrected too -- and refuses a review that leaves an
amber figure unanswered. What it answers is one period per scope
(standalone, consolidated) in the compact ``FORMAT`` that
``coverage_periods.figures`` holds: every figure with its page, and what the
analyst did with it.

``workbook`` writes a stored period as a spreadsheet: one sheet a
statement, the subtotals the statement states written as **formulas** of
the figures they add up (found again on the accepted figures by the
reader's own sums, ``checks.identities``), each figure's page and what the
analyst did in a note. A formula is written only where the parts add to the
printed total exactly, so the workbook recalculates to the totals stored;
a printed total its parts do not reach stays a figure and its note says so.

``stored_for_read`` gives what is kept for a company in the shape the
reader's comparison with stored figures takes (``checks.against_stored``),
converted to the unit the new filing prints.

Pure: no database, no provider, no clock but the one passed in.
"""
from __future__ import annotations

import datetime
import io
import math
from typing import Any, Iterable, Mapping, Optional, Sequence

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.money import UNIT_SCALE
from research.reader import checks
from research.workbook import _put

FORMAT = 1
SCOPES = ("consolidated", "standalone")
UNITS = ("units", "thousands", "lakhs", "millions", "crores", "billions")
SCALE = {"units": 1.0, **UNIT_SCALE}
MAX_FIGURE = 1e15
MAX_CORRECTIONS = 5000

# What the analyst did with a figure (a mark beside it; none = the reads agreed and nothing was doubtful)
ACCEPTED = "accepted"
CORRECTED = "corrected"

# Refusals, as codes the screen words (web/messages/en.json coverage.error_*)
AMBER_UNREVIEWED = "amber_unreviewed"
BAD_REVIEW = "bad_review"
QUARTER_UNKNOWN = "quarter_unknown"
UNIT_UNKNOWN = "unit_unknown"
NO_FIGURES = "no_figures"

STATEMENT_WORDS = {"results": "Results", "segments": "Segments", "balance_sheet": "Balance sheet",
                   "cash_flow": "Cash flow"}
AMBER_WORDS = {
    checks.READS_DIFFER: "the two reads gave it differently",
    checks.SIGN_DIFFERS: "the two reads gave it with different signs",
    checks.ONE_READ_ONLY: "only one read had this row",
    checks.SUM_FAILS: "a sum the statement states did not add up here",
    checks.POOR_SCAN: "it is on a poorly scanned page",
    checks.UNBALANCED: "the balance sheet did not balance in this column",
    checks.CASH_MISMATCH: "the change in cash did not reach the closing cash",
    checks.STORED_DIFFERS: "it differed from the figure already kept for this period",
}


class Refused(ValueError):
    """The review can't be accepted; ``code`` says why (and ``count`` how many figures)."""

    def __init__(self, code: str, count: int = 0):
        super().__init__(code)
        self.code, self.count = code, count


def _figure(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) \
            or abs(value) > MAX_FIGURE:
        raise Refused(BAD_REVIEW)
    return float(value)


def _cell(result: Mapping, ref: Any) -> tuple[int, int, int]:
    """A (table, row, column) the review names, checked against the read."""
    try:
        t, r, c = (int(x) for x in ref)
        cells = result["tables"][t]["rows"][r]["cells"]
        cells[c]
    except (TypeError, ValueError, IndexError, KeyError):
        raise Refused(BAD_REVIEW) from None
    if min(t, r, c) < 0:
        raise Refused(BAD_REVIEW)
    return t, r, c


def amber_cells(result: Mapping) -> set[tuple[int, int, int]]:
    """Every figure of the read the analyst must answer: amber, with a value."""
    return {(t, r, c) for t, table in enumerate(result["tables"]) for r, row in enumerate(table["rows"])
            for c, cell in enumerate(row["cells"]) if cell.get("amber") and cell.get("value") is not None}


def period_months(result: Mapping) -> int:
    """What the filing is for: a quarter when its results print a three-month
    column ending at the quarter end, else the shortest period ending there."""
    end = result.get("quarter_end")
    found = sorted({c.get("months") for t in result["tables"] if t["statement"] in ("results", "segments")
                    for c in t["columns"] if c.get("period_end") == end and c.get("months") in (3, 6, 9, 12)})
    return found[0] if found else 3


def accept(result: Mapping, review: Mapping, *, unit: Optional[str] = None) -> list[dict]:
    """The periods a reviewed read gives, one per scope: ``[{period_end,
    months, scope, unit, currency, figures}]``. ``review`` holds
    ``accepted``: amber figures kept as read, as ``[table, row, column]``,
    and ``corrected``: ``{table, row, column, value}`` (value null for a
    blank). Raises ``Refused``."""
    if not result.get("quarter_end"):
        raise Refused(QUARTER_UNKNOWN)
    unit = result.get("unit") or unit
    if unit not in UNITS:
        raise Refused(UNIT_UNKNOWN)
    corrections = review.get("corrected") or []
    accepted = review.get("accepted") or []
    if len(corrections) + len(accepted) > MAX_CORRECTIONS:
        raise Refused(BAD_REVIEW)
    fixed: dict[tuple, Optional[float]] = {}
    for item in corrections:
        if not isinstance(item, Mapping):
            raise Refused(BAD_REVIEW)
        fixed[_cell(result, (item.get("table"), item.get("row"), item.get("column")))] = _figure(item.get("value"))
    kept = {_cell(result, ref) for ref in accepted}
    unanswered = amber_cells(result) - kept - set(fixed)
    if unanswered:
        raise Refused(AMBER_UNREVIEWED, len(unanswered))

    months = period_months(result)
    by_scope: dict[str, list] = {}
    for t, table in enumerate(result["tables"]):
        rows = []
        for r, row in enumerate(table["rows"]):
            values, marks = [], []
            for c, cell in enumerate(row["cells"]):
                read = cell.get("value")
                if (t, r, c) in fixed:
                    value = fixed[(t, r, c)]
                    changed = value != read
                    mark = {"do": CORRECTED, "read": read, "amber": cell.get("amber")} if changed else (
                        {"do": ACCEPTED, "amber": cell["amber"]} if cell.get("amber") else None)
                else:
                    value = read
                    mark = {"do": ACCEPTED, "amber": cell["amber"]} if cell.get("amber") and read is not None else None
                values.append(value)
                marks.append(mark)
            rows.append({"label": row["label"], "line": row.get("line"), "page": row.get("page"),
                         "values": values, **({"marks": marks} if any(marks) else {})})
        by_scope.setdefault(table["scope"], []).append({
            "statement": table["statement"], "pages": list(table.get("pages") or []),
            "columns": [{k: col.get(k) for k in ("label", "period_end", "months", "audited")}
                        for col in table["columns"]],
            "rows": rows})
    if not by_scope:
        raise Refused(NO_FIGURES)
    document = result.get("document") or {}
    out = []
    for scope in SCOPES:
        if scope not in by_scope:
            continue
        out.append({"period_end": result["quarter_end"], "months": months, "scope": scope, "unit": unit,
                    "currency": result.get("currency") or None,
                    "figures": {"format": FORMAT, "company": result.get("company") or "",
                                "year_end_month": result.get("year_end_month"),
                                "document_pages": document.get("pages"), "tables": by_scope[scope]}})
    return out


def counts(figures: Mapping) -> dict:
    """How many figures a stored period holds (as the reader counts them: a
    zero dash is none), and how many the analyst corrected or accepted."""
    total = corrected = accepted = 0
    for table in figures.get("tables") or []:
        for row in table["rows"]:
            total += sum(v is not None and v != 0 for v in row["values"])
            for mark in row.get("marks") or []:
                if mark:
                    corrected += mark["do"] == CORRECTED
                    accepted += mark["do"] == ACCEPTED
    return {"figures": total, "corrected": corrected, "accepted": accepted}


def stored_for_read(periods: Iterable[Mapping], unit: Optional[str]) -> list[dict]:
    """What is kept for a company, figure by figure on its standard lines, in
    ``unit`` (the new filing's): the shape ``checks.against_stored`` takes."""
    if unit not in SCALE:
        return []
    out = []
    for period in periods:
        factor = SCALE[period["unit"]] / SCALE[unit]
        for table in period["figures"].get("tables") or []:
            for row in table["rows"]:
                if not row.get("line"):
                    continue
                for column, value in zip(table["columns"], row["values"]):
                    if value is None:
                        continue
                    out.append({"statement": table["statement"], "scope": period["scope"], "line": row["line"],
                                "period_end": column.get("period_end"), "months": column.get("months"),
                                "value": round(value * factor, 6)})
    return out


# ---------------------------------------------------------------------------
# The checked workbook
# ---------------------------------------------------------------------------
FIRST_ROW = 6          # the first figure's row; above it the title, the unit and two heading rows
FIRST_COLUMN = 3       # A: the company's label, B: the standard line


def _sheet_title(name: str, used: set) -> str:
    title, n = name[:31], 2
    while title.lower() in used:
        title, n = f"{name[:28]} {n}", n + 1
    used.add(title.lower())
    return title


def _decimals(rows: Sequence[Sequence[Optional[float]]]) -> int:
    return 2 if checks.step(rows) < 1 else 0


def subtotals(table: Mapping) -> tuple[dict, dict]:
    """The statement's own sums on the accepted figures, in one search:
    those that add up exactly, ``{(row, column): [(part row, sign), ...]}``,
    and the printed totals they don't reach, ``{(row, column): what the
    parts add to}``."""
    grid = [list(row["values"]) for row in table["rows"]]
    found = checks.identities({"columns": table["columns"],
                               "rows": [{"label": row["label"], "values": row["values"]} for row in table["rows"]]})
    decimals = _decimals(grid)
    exact, off = {}, {}
    for item in found["found"]:
        target = item["row"]
        for column in range(len(table["columns"])):
            printed = grid[target][column]
            if printed is None:
                continue
            parts = [(p, s) for p, s in zip(item["parts"], item["signs"]) if grid[p][column] is not None]
            worked = sum(s * grid[p][column] for p, s in parts)
            if len(parts) >= 2 and round(worked, decimals) == round(printed, decimals) \
                    and abs(worked - printed) < 0.5 * 10 ** -decimals:
                exact[(target, column)] = parts
            else:
                off[(target, column)] = round(worked, decimals)
    return exact, off


def _formula(parts: list[tuple[int, int]], column: int) -> str:
    letter = get_column_letter(FIRST_COLUMN + column)
    text = ""
    for row, sign in parts:
        ref = f"{letter}{FIRST_ROW + row}"
        text += (("+" if text else "") if sign > 0 else "-") + ref
    return "=" + text


def _note(row: Mapping, column: int, sum_parts: Optional[list], printed_off: Optional[float]) -> str:
    lines = [f"Page {row['page']}" if row.get("page") else "Page not known"]
    mark = (row.get("marks") or [None] * (column + 1))[column]
    if mark and mark["do"] == CORRECTED:
        read = mark.get("read")
        lines.append("Corrected by the analyst; the reader read " + ("a blank" if read is None else f"{read:,}"))
    elif mark and mark["do"] == ACCEPTED:
        lines.append("Checked and accepted as read by the analyst: " + AMBER_WORDS.get(mark.get("amber"), "it was flagged"))
    if sum_parts:
        lines.append("A subtotal: a formula of the figures it adds up")
    elif printed_off is not None:
        lines.append(f"The printed total; its parts add to {printed_off:,}")
    return "\n".join(lines)


COLUMN_WIDTH = 16


def layout(company_name: str, period: Mapping, *, generated: datetime.datetime, source: str) -> list[dict]:
    """The stored period as sheets of cells, the one description both the
    download (``workbook``) and the Excel add-in (PLAN.md R3, ``GET
    /api/research/periods/{id}/cells``) write, so the two can never hold
    different cells: an About sheet, then a sheet a statement.

    A sheet is ``{"title", "cells", "widths", "freeze"}``; a cell is
    ``{"row", "column"}`` (from 1) with exactly one of ``text`` (words, never
    a formula, whatever they start with), ``number`` or ``formula``, and
    optionally ``number_format``, ``bold`` and ``note``. ``widths`` maps a
    column letter to its width in characters; ``freeze`` is the first cell
    that scrolls (``"C6"``) or None."""
    figures = period["figures"]
    scope = period["scope"]
    unit_words = "as printed" if period["unit"] == "units" else period["unit"]
    currency = period.get("currency") or "currency not printed"
    items = [("Source", source), ("Generated (UTC)", generated.strftime("%Y-%m-%d %H:%M:%S")),
             ("Company", company_name), ("Name in the filing", figures.get("company") or ""),
             ("Period ends", str(period["period_end"])), ("Months", str(period["months"])),
             ("Scope", scope), ("Currency", currency), ("Money unit", unit_words),
             ("Accepted (UTC)", str(period.get("accepted_at") or "")),
             ("How it was made", "Read twice from the results filing by the results reader, checked, "
                                 "and reviewed and accepted by the analyst. Subtotals are formulas of "
                                 "the figures they add up; each figure's note gives its page.")]
    about = [_text(1, 1, "Item", bold=True), _text(1, 2, "Value", bold=True)]
    for i, (item, value) in enumerate(items, start=2):
        about += [_text(i, 1, item), _text(i, 2, value)]
    sheets = [{"title": "About", "cells": about, "widths": {"A": 22, "B": 80}, "freeze": None}]
    used = {"about"}
    for table in figures.get("tables") or []:
        title = _sheet_title(f"{STATEMENT_WORDS.get(table['statement'], table['statement'])} {scope}", used)
        sheets.append({"title": title, **_table_sheet(company_name, table, scope, unit_words, currency)})
    return sheets


def _text(row: int, column: int, text: str, *, bold: bool = False) -> dict:
    cell = {"row": row, "column": column, "text": text}
    if bold:
        cell["bold"] = True
    return cell


def _table_sheet(company_name: str, table: Mapping, scope: str, unit_words: str, currency: str) -> dict:
    grid = [list(row["values"]) for row in table["rows"]]
    number_format = "#,##0.00" if _decimals(grid) else "#,##0"
    sums, worked_off = subtotals(table)
    statement = STATEMENT_WORDS.get(table["statement"], table["statement"])
    cells = [_text(1, 1, f"{company_name}: {statement}, {scope}", bold=True),
             _text(2, 1, f"{currency}, {unit_words}. Pages "
                         f"{', '.join(str(p) for p in table.get('pages') or []) or 'not known'}."),
             _text(4, 1, "Line as printed", bold=True), _text(4, 2, "Standard line", bold=True),
             _text(5, 1, "Period ends (months)")]
    for c, column in enumerate(table["columns"]):
        months = column.get("months")
        cells += [_text(4, FIRST_COLUMN + c, column.get("label") or "", bold=True),
                  _text(5, FIRST_COLUMN + c, f"{column.get('period_end')} ({months if months is not None else '-'})")]
    for r, row in enumerate(table["rows"]):
        at = FIRST_ROW + r
        cells.append(_text(at, 1, row["label"]))
        if row.get("line"):
            cells.append(_text(at, 2, row["line"]))
        for c, value in enumerate(row["values"]):
            if value is None:
                continue
            parts = sums.get((r, c))
            cell = {"row": at, "column": FIRST_COLUMN + c, "number_format": number_format,
                    "note": _note(row, c, parts, worked_off.get((r, c)))}
            if parts:
                cell["formula"] = _formula(parts, c)
            else:
                cell["number"] = value
            cells.append(cell)
    widths = {"A": 48, "B": 26, **{get_column_letter(FIRST_COLUMN + c): COLUMN_WIDTH for c in range(len(table["columns"]))}}
    return {"cells": cells, "widths": widths, "freeze": f"{get_column_letter(FIRST_COLUMN)}{FIRST_ROW}"}


def workbook(company_name: str, period: Mapping, *, generated: datetime.datetime, source: str) -> Workbook:
    """The stored period as a workbook, written from ``layout``."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in layout(company_name, period, generated=generated, source=source):
        ws = wb.create_sheet(sheet["title"])
        for spec in sheet["cells"]:
            if "text" in spec:
                cell = _put(ws, spec["row"], spec["column"], spec["text"])
            else:
                cell = ws.cell(spec["row"], spec["column"])
                cell.value = spec["formula"] if "formula" in spec else spec["number"]
            if spec.get("number_format"):
                cell.number_format = spec["number_format"]
            if spec.get("bold"):
                cell.font = Font(bold=True)
            if spec.get("note"):
                cell.comment = Comment(spec["note"], "variater.com", width=260, height=90)
        for letter, width in sheet["widths"].items():
            ws.column_dimensions[letter].width = width
        if sheet["freeze"]:
            ws.freeze_panes = sheet["freeze"]
    wb.properties.creator = source
    wb.properties.lastModifiedBy = source
    return wb


def workbook_bytes(company_name: str, period: Mapping, *, generated: datetime.datetime, source: str) -> bytes:
    buf = io.BytesIO()
    workbook(company_name, period, generated=generated, source=source).save(buf)
    return buf.getvalue()
