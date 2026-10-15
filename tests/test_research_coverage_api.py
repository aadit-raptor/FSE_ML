"""Coverage through the API (PLAN.md R2): a read accepted into a company,
corrected figures stored and downloaded, another account shut out, and the
limits.

The road is the real one -- upload, scan, read (recorded answers, no AI),
review, accept, download -- on a throwaway database
(tests/test_research_reader_api.py's fixtures).
"""
from __future__ import annotations

import io
import json
import logging

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.main import app
from db import coverage as rows
from research import coverage
from tests.research_support import workbook_values
from tests.test_research_reader_api import clean_pdf, reader, run_reads  # noqa: F401 -- fixtures
from tests.test_uploads import uploads  # noqa: F401 -- the fixture

client = TestClient(app)
AVENUE = "avenue-supermarts_FY27-Q2"


@pytest.fixture
def avenue(reader, monkeypatch):  # noqa: F811 -- the imported fixture
    monkeypatch.setenv("FSE_READER_REPLAY_DEFAULT", AVENUE)
    return reader


def read(queue, file_id: str, **extra) -> dict:
    started = client.post("/api/research/reads", json={"file_id": file_id, "public": True,
                                                       "quarter_end": "2026-09-30", **extra})
    assert started.status_code == 202, started.text
    run_reads(queue)
    done = client.get(f"/api/research/reads/{started.json()['id']}").json()
    assert done["status"] == "succeeded", done
    return done


def company(name: str = "DMart") -> str:
    made = client.post("/api/research/companies", json={"name": name})
    assert made.status_code == 201, made.text
    return made.json()["id"]


def amber(result: dict) -> list:
    return [[t, r, c] for t, table in enumerate(result["tables"]) for r, row in enumerate(table["rows"])
            for c, cell in enumerate(row["cells"]) if cell["amber"] and cell["value"] is not None]


def accept(company_id: str, read_id: str, **review):
    return client.post(f"/api/research/companies/{company_id}/periods", json={"read_id": read_id, **review})


def test_a_read_becomes_an_accepted_quarter_only_when_the_analyst_accepts_it(avenue, job_queue):
    file_id = clean_pdf(job_queue)
    done = read(job_queue, file_id)
    firm = company()
    # nothing is stored by the read itself
    assert client.get("/api/research/companies").json()["companies"][0]["periods"] == []
    # nor while an amber figure is unanswered
    refused = accept(firm, done["id"])
    assert (refused.status_code, refused.json()["detail"]) == (422, "amber_unreviewed")
    (t, r, c), = amber(done["result"])
    made = accept(firm, done["id"], corrected=[{"table": t, "row": r, "column": c, "value": 4321.5}])
    assert made.status_code == 201, made.text
    period, = made.json()["periods"]
    assert (period["period_end"], period["months"], period["scope"], period["unit"], period["currency"],
            period["file_id"]) == ("2026-09-30", 3, "consolidated", "crores", "INR", file_id)
    listed = client.get("/api/research/companies").json()
    assert [p["id"] for p in listed["companies"][0]["periods"]] == [period["id"]]
    assert listed["limits"] == {"companies": rows.MAX_COMPANIES, "periods": rows.MAX_PERIODS}

    # the corrected figure is the one stored ...
    kept = client.get(f"/api/research/periods/{period['id']}").json()
    row = kept["figures"]["tables"][t]["rows"][r]
    assert row["values"][c] == 4321.5 and row["marks"][c]["do"] == "corrected"
    assert kept["counts"]["corrected"] == 1 and kept["company_name"] == "DMart"
    # ... and the one downloaded, in a workbook that recalculates to every stored figure
    got = client.get(f"/api/research/periods/{period['id']}/workbook")
    assert got.status_code == 200
    assert got.headers["content-disposition"] == 'attachment; filename="DMart_2026-09-30_consolidated.xlsx"'
    wb = load_workbook(io.BytesIO(got.content))
    sheets = wb.sheetnames[1:]
    for sheet, table in zip(sheets, kept["figures"]["tables"]):
        worked = workbook_values(wb, sheet)
        for i, stored_row in enumerate(table["rows"]):
            for j, value in enumerate(stored_row["values"]):
                if value is not None:
                    address = wb[sheet].cell(coverage.FIRST_ROW + i, coverage.FIRST_COLUMN + j).coordinate
                    assert worked(address) == pytest.approx(value, abs=1e-6), (sheet, address)
    assert wb[sheets[t]].cell(coverage.FIRST_ROW + r, coverage.FIRST_COLUMN + c).value == 4321.5
    # each step is in the account's history, without a name or a figure
    actions = [e["action"] for e in client.get("/api/account/history").json()["entries"]]
    assert {"coverage_added", "period_accepted", "exported"} <= set(actions)


def test_accepting_the_same_quarter_again_replaces_it(avenue, job_queue):
    file_id = clean_pdf(job_queue)
    done = read(job_queue, file_id)
    firm = company()
    first = accept(firm, done["id"], accepted=amber(done["result"])).json()["periods"][0]
    (t, r, c), = amber(done["result"])
    second = accept(firm, done["id"], corrected=[{"table": t, "row": r, "column": c, "value": 1.0}]).json()
    assert second["periods"][0]["id"] == first["id"]
    assert len(client.get("/api/research/companies").json()["companies"][0]["periods"]) == 1
    kept = client.get(f"/api/research/periods/{first['id']}").json()
    assert kept["figures"]["tables"][t]["rows"][r]["values"][c] == 1.0


def test_another_account_cannot_read_change_or_use_anything(avenue, job_queue):
    file_id = clean_pdf(job_queue)
    done = read(job_queue, file_id)
    firm = company()
    period = accept(firm, done["id"], accepted=amber(done["result"])).json()["periods"][0]["id"]
    avenue("test:mallory")
    assert client.get("/api/research/companies").json()["companies"] == []
    for method, path in (("get", f"/api/research/periods/{period}"),
                         ("get", f"/api/research/periods/{period}/workbook"),
                         ("get", f"/api/research/periods/{period}/cells"),
                         ("delete", f"/api/research/periods/{period}"),
                         ("delete", f"/api/research/companies/{firm}"),
                         ("get", f"/api/files/{file_id}/content")):
        assert getattr(client, method)(path).status_code == 404, path
    # nor accept the owner's read into the owner's company, or into their own
    assert accept(firm, done["id"], accepted=amber(done["result"])).status_code == 404
    own = company("Mine")
    assert accept(own, done["id"], accepted=amber(done["result"])).json()["detail"] == "no_read"
    # nor read into the owner's company
    refused = client.post("/api/research/reads", json={"file_id": file_id, "public": True, "company_id": firm})
    assert refused.status_code == 404
    avenue("test:alice")
    assert client.get(f"/api/research/periods/{period}").status_code == 200


def test_the_source_pdf_is_handed_to_its_owner_for_the_page_view(avenue, job_queue):
    from tests import upload_samples as s

    file_id = clean_pdf(job_queue)
    got = client.get(f"/api/files/{file_id}/content")
    assert got.status_code == 200 and got.content == s.pdf()
    assert got.headers["content-type"] == "application/pdf"
    assert got.headers["content-disposition"] == "attachment"


def test_a_read_for_a_covered_company_is_checked_against_what_is_kept(avenue, job_queue):
    file_id = clean_pdf(job_queue)
    firm = company()
    done = read(job_queue, file_id, company_id=firm)
    (t, r, c), = amber(done["result"])
    revenue = next(i for i, row in enumerate(done["result"]["tables"][0]["rows"])
                   if row["line"] == "revenue_from_operations")
    accept(firm, done["id"], accepted=[[t, r, c]],
           corrected=[{"table": 0, "row": revenue, "column": 0, "value": 20000.0}])
    again = read(job_queue, file_id, company_id=firm)
    cell = again["result"]["tables"][0]["rows"][revenue]["cells"][0]
    assert (cell["amber"], cell["other"]) == ("stored_differs", 20000.0)


def test_unfinished_and_unknown_reads_are_refused(avenue, job_queue):
    firm = company()
    file_id = clean_pdf(job_queue)
    started = client.post("/api/research/reads", json={"file_id": file_id, "public": True}).json()
    assert accept(firm, started["id"]).json()["detail"] == "read_not_done"
    record = job_queue.submit("test:alice", "montecarlo.run", {}, max_active_per_owner=4, max_queued=10)
    assert accept(firm, str(record.id)).json()["detail"] == "no_read"


def test_a_person_keeps_at_most_their_companies_and_periods(avenue, job_queue, monkeypatch):
    monkeypatch.setattr(rows, "MAX_COMPANIES", 1)
    firm = company()
    refused = client.post("/api/research/companies", json={"name": "Two"})
    assert (refused.status_code, refused.json()["detail"]) == (409, "too_many_companies")
    file_id = clean_pdf(job_queue)
    done = read(job_queue, file_id)
    monkeypatch.setattr(rows, "MAX_PERIODS", 0)
    full = accept(firm, done["id"], accepted=amber(done["result"]))
    assert (full.status_code, full.json()["detail"]) == (409, "too_many_periods")
    monkeypatch.setattr(rows, "MAX_PERIODS", 10)
    monkeypatch.setattr(rows, "BUDGET_BYTES", 1)
    assert accept(firm, done["id"], accepted=amber(done["result"])).json()["detail"] == "coverage_full"


def test_removing_a_company_removes_its_periods_and_a_deleted_file_leaves_the_figures(avenue, job_queue):
    file_id = clean_pdf(job_queue)
    done = read(job_queue, file_id)
    firm = company()
    period = accept(firm, done["id"], accepted=amber(done["result"])).json()["periods"][0]["id"]
    assert client.delete(f"/api/files/{file_id}").status_code == 204
    kept = client.get(f"/api/research/periods/{period}").json()
    assert kept["file_id"] is None and kept["counts"]["figures"] == 332
    assert client.delete(f"/api/research/companies/{firm}").status_code == 204
    assert client.get(f"/api/research/periods/{period}").status_code == 404


def test_nothing_from_the_filing_or_the_name_reaches_the_log(avenue, job_queue, caplog):
    caplog.set_level(logging.DEBUG)
    file_id = clean_pdf(job_queue)
    done = read(job_queue, file_id)
    firm = company("Secret Retail Co")
    period = accept(firm, done["id"], accepted=amber(done["result"])).json()["periods"][0]["id"]
    client.get(f"/api/research/periods/{period}/workbook")
    logged = "\n".join(r.getMessage() for r in caplog.records) + json.dumps([r.__dict__ for r in caplog.records],
                                                                          default=str)
    for secret in ("Secret Retail Co", "Avenue Supermarts", "19644.01", "19,644.01"):
        assert secret not in logged


def test_pdf_views_are_counted_against_a_daily_limit(avenue, job_queue, monkeypatch):
    from api.limits import USER_FILE_VIEWS, Rule

    file_id = clean_pdf(job_queue)
    monkeypatch.setattr("api.routers.files.USER_FILE_VIEWS",
                        Rule(USER_FILE_VIEWS.name, "PDF views", "user", 2, 86_400, shared=True))
    assert [client.get(f"/api/files/{file_id}/content").status_code for _ in range(3)] == [200, 200, 429]


def test_a_pdf_view_comes_from_the_store_once_the_upload_spool_is_gone(avenue, job_queue, tmp_path):
    from tests import upload_samples as s
    from uploads import service

    file_id = clean_pdf(job_queue)
    service.spooled(file_id).unlink(missing_ok=True)
    got = client.get(f"/api/files/{file_id}/content")
    assert got.status_code == 200 and got.content == s.pdf()
    # the store's copy, fetched to the spool for the answer, is not left behind
    assert not list((tmp_path / "spool").glob("view-*"))


# ------------------------------------------------------------------ the Excel add-in's cells (PLAN.md R3)

def as_written(answer: dict) -> dict:
    """The cells answer in the terms ``tests.addin_cells_check`` compares: what the pane asks Excel to hold."""
    from openpyxl.utils import get_column_letter
    sheets = []
    for sheet in answer["sheets"]:
        cells = {}
        for cell in sheet["cells"]:
            spec = {"formula": cell["formula"]} if cell.get("formula") else \
                {"value": cell["text"] if cell.get("text") is not None else cell["number"]}
            if cell.get("number_format"):
                spec["format"] = cell["number_format"]
            if cell.get("bold"):
                spec["bold"] = True
            if cell.get("note"):
                spec["note"] = cell["note"]
            cells[f"{get_column_letter(cell['column'])}{cell['row']}"] = spec
        sheets.append({"name": sheet["title"], "cells": cells, "widths": sheet["widths"], "freeze": sheet["freeze"]})
    return {"sheets": sheets}


def test_the_add_in_is_given_the_same_cells_as_the_download(avenue, job_queue):
    from tests.addin_cells_check import differences
    done = read(job_queue, clean_pdf(job_queue))
    firm = company()
    (t, r, c), = amber(done["result"])
    period = accept(firm, done["id"], corrected=[{"table": t, "row": r, "column": c, "value": 4321.5}]
                    ).json()["periods"][0]["id"]
    answer = client.get(f"/api/research/periods/{period}/cells")
    assert answer.status_code == 200, answer.text
    body = answer.json()
    assert (body["period_id"], body["company_name"]) == (period, "DMart")
    wb = load_workbook(io.BytesIO(client.get(f"/api/research/periods/{period}/workbook").content))
    problems, count = differences(wb, as_written(body))
    assert problems == []
    assert count > 300
    # subtotals arrive as formulas, the corrected figure as the analyst's
    formulas = [cell for sheet in body["sheets"] for cell in sheet["cells"] if cell.get("formula")]
    assert len(formulas) > 40
    assert any(cell.get("number") == 4321.5 and "Corrected by the analyst" in cell["note"]
               for sheet in body["sheets"] for cell in sheet["cells"])
    # and the comparison itself notices a changed cell
    changed = as_written(body)
    first = next(spec for spec in changed["sheets"][1]["cells"].values() if "formula" in spec)
    first["formula"] += "+1"
    assert differences(wb, changed)[0]


def test_writing_into_excel_is_recorded_as_an_export(avenue, job_queue):
    done = read(job_queue, clean_pdf(job_queue))
    period = accept(company(), done["id"], accepted=amber(done["result"])).json()["periods"][0]["id"]
    assert client.get(f"/api/research/periods/{period}/cells").status_code == 200
    history = client.get("/api/account/history").json()
    assert [e.get("export") for e in history["entries"] if e["action"] == "exported"] == ["results_excel"]
