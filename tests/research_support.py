"""Shared by the research tests (PLAN.md R4 part b).

Three things the tests need and the product does not:

- the two hand-solved companies of model 1 (`tests/fixtures/research/`: the
  figures typed from their filings, which lines make each model row, and what
  the hand-built Excel workbooks gave), turned into a `Company`;
- **an independent reader of a written workbook**: its own small reader of
  Excel formula text, which shares nothing with `research/formula.py`, so a
  mistake in how the writer turns a formula into Excel shows as a difference;
- LibreOffice working a workbook out, where it is installed.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from research import formula
from research.definition import Company

FIXTURES = Path(__file__).parent / "fixtures" / "research"
PAST = ("ly", "ytd")


def load(slug: str) -> dict:
    return json.loads((FIXTURES / ("model1_%s.json" % slug)).read_text(encoding="utf-8"))


def reported(fixture: dict) -> dict:
    """{reported key: {column label: figure}}: the typed figures, their own workings worked out."""
    out = {}
    for section in fixture["sections"]:
        for j, column in enumerate(section["columns"]):
            names = {}
            for row in section["rows"]:
                v = row["values"][j]
                if v is None:
                    continue
                names[row["key"]] = formula.value(v, names) if isinstance(v, str) else v
                out.setdefault(row["key"], {})[column] = names[row["key"]]
    return out


def _section_of(fixture: dict, key: str) -> dict:
    return next(s for s in fixture["sections"] if any(r["key"] == key for r in s["rows"]))


def past_figures(fixture: dict) -> tuple[dict, dict]:
    """The model's typed past figures, and the marks they inherit, from the reported lines.

    This is the arithmetic the results reader will do one day (PLAN.md R2):
    finance cost less interest on leases, the fixed asset lines added up.
    """
    figures, values, marks = reported(fixture), {}, {}
    for row, text in fixture["map"].items():
        used = sorted({ref.name for ref in formula.refs(formula.parse(text))})
        for period in PAST:
            columns = {key: _section_of(fixture, key)["use"].get(period) for key in used}
            if any(c is None or c not in figures.get(k, {}) for k, c in columns.items()):
                continue
            values.setdefault(row, {})[period] = formula.value(text, {k: figures[k][c] for k, c in columns.items()})
            for key, column in columns.items():
                section = _section_of(fixture, key)
                line = next(r for r in section["rows"] if r["key"] == key)
                level = (line.get("marks") or [None] * len(section["columns"]))[section["columns"].index(column)]
                if level:
                    was = marks.setdefault(row, {}).get(period)
                    marks[row][period] = {"level": "red" if "red" in (level, was and was["level"]) else "amber", "why": line["why"]}
    return values, marks


def company(fixture: dict, **changes) -> Company:
    """The fixture as the model's input. `changes` are `row={period: value}` overrides."""
    values, marks = past_figures(fixture)
    for row, by_period in fixture["assumptions"].items():
        values.setdefault(row, {}).update(by_period)
    for row, by_period in changes.items():
        values.setdefault(row, {}).update(by_period)
    return Company(
        name=fixture["name"], sector=fixture["sector"], unit_name=fixture["unit_name"], basis=fixture["basis"],
        words=fixture["words"], period_labels=fixture["period_labels"],
        constants={"unit_rupees": fixture["unit_rupees"]}, choices=fixture["choices"],
        values=values, off=tuple(fixture["off"]), marks=marks)


# ---------------------------------------------------------------------------
# An independent reader of Excel formulas: numbers, cells, + - * /, brackets,
# MAX, MIN and IF(a=b, x, y). Anything else is an error, which is itself a
# test: the workbook may hold nothing more.
# ---------------------------------------------------------------------------
_XL = re.compile(r"\s*(?:(?P<cell>(?:[A-Za-z_]+!)?\$?[A-Z]{1,3}\$?\d+)|(?P<num>\d+(?:\.\d+)?(?:[eE][-+]?\d+)?)|(?P<fn>[A-Z]+)(?=\()|(?P<op>[-+*/(),=]))")


class _Excel:
    def __init__(self, text: str, cell):
        self.text, self.cell, self.at = text, cell, 0
        self.tokens = []
        pos = 0
        while pos < len(text):
            m = _XL.match(text, pos)
            if m is None:
                raise ValueError("the workbook holds something the reader does not know: %r" % text[pos:pos + 20])
            self.tokens.append((m.lastgroup, m.group(m.lastgroup)))
            pos = m.end()

    def peek(self):
        return self.tokens[self.at][1] if self.at < len(self.tokens) else None

    def take(self, want=None):
        kind, v = self.tokens[self.at]
        assert want is None or v == want, "expected %r, found %r in %r" % (want, v, self.text)
        self.at += 1
        return kind, v

    def skip(self):
        """Pass over one argument without working it out (the branch IF does not take)."""
        depth = 0
        while True:
            v = self.peek()
            if depth == 0 and v in (",", ")"):
                return
            depth += (v == "(") - (v == ")")
            self.at += 1

    def expression(self):
        v = self.term()
        while self.peek() in ("+", "-"):
            v = v + self.term() if self.take()[1] == "+" else v - self.term()
        return v

    def term(self):
        v = self.unary()
        while self.peek() in ("*", "/"):
            v = v * self.unary() if self.take()[1] == "*" else v / self.unary()
        return v

    def unary(self):
        if self.peek() == "-":
            self.take()
            return -self.unary()
        kind, v = self.take()
        if kind == "num":
            return float(v)
        if kind == "cell":
            return self.cell(v.replace("$", ""))
        if kind == "fn":
            return self.function(v)
        assert v == "(", "unexpected %r in %r" % (v, self.text)
        inner = self.expression()
        self.take(")")
        return inner

    def function(self, name):
        self.take("(")
        if name == "IF":
            left = self.expression()
            self.take("=")
            holds = left == self.expression()
            self.take(",")
            if holds:
                v = self.expression()
                self.take(",")
                self.skip()
            else:
                self.skip()
                self.take(",")
                v = self.expression()
            self.take(")")
            return v
        assert name in ("MAX", "MIN"), "the workbook uses %s" % name
        args = [self.expression()]
        while self.peek() == ",":
            self.take()
            args.append(self.expression())
        self.take(")")
        return max(args) if name == "MAX" else min(args)


def workbook_values(wb, sheet: str = "Model"):
    """A function giving any cell's value, worked out from the workbook's own formulas.

    An empty cell is None; a formula naming an empty cell reads it as 0, as
    Excel does.
    """
    cache = {}

    def get(address: str):
        name, _, a = address.rpartition("!")
        key = (name or sheet, a)
        if key not in cache:
            raw = wb[key[0]][a].value
            if isinstance(raw, str) and raw.startswith("="):
                reader = _Excel(raw[1:], lambda ref: get(ref if "!" in ref else key[0] + "!" + ref) or 0.0)
                raw = reader.expression()
                assert reader.at == len(reader.tokens), "left over in %r" % wb[key[0]][a].value
            cache[key] = raw
        return cache[key]

    return get


# ---------------------------------------------------------------------------
# A real spreadsheet engine: LibreOffice, without a screen.
# ---------------------------------------------------------------------------
def soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def soffice_wanted() -> bool:
    """True where the LibreOffice test must run: it is installed, or CI says it has to be."""
    return soffice() is not None or os.environ.get("FSE_REQUIRE_SOFFICE") == "1"


# LibreOffice trusts the values stored in a workbook that says Excel wrote it
# and ours stores none, so tell it to work every formula out on loading.
_ALWAYS_RECALCULATE = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def recalculated_by_libreoffice(wb):
    """The workbook after LibreOffice has opened it, worked it out and saved it (values only)."""
    binary = soffice()
    assert binary, "LibreOffice is required here (FSE_REQUIRE_SOFFICE=1) and `soffice` is not on the PATH"
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile" / "user"
        profile.mkdir(parents=True)
        (profile / "registrymodifications.xcu").write_text(_ALWAYS_RECALCULATE, encoding="utf-8")
        source = Path(tmp) / "in" / "model.xlsx"
        source.parent.mkdir()
        wb.save(source)
        out = Path(tmp) / "out"
        done = subprocess.run(
            [binary, "-env:UserInstallation=" + (Path(tmp) / "profile").as_uri(), "--headless", "--norestore",
             "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(source)],
            capture_output=True, text=True, timeout=180)
        result = out / "model.xlsx"
        assert result.exists(), "LibreOffice wrote nothing: %s %s" % (done.stdout[-400:], done.stderr[-400:])
        return load_workbook(result, data_only=True)
