"""Excel downloads: the workbook holds exactly the tables sent, and the Monte
Carlo sample is the simulation behind the on-screen results."""
import io

import pandas as pd
import pytest
from fastapi.testclient import TestClient

from api.main import app

client = TestClient(app)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _read(resp):
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == XLSX
    return pd.read_excel(io.BytesIO(resp.content), sheet_name=None)


def test_workbook_round_trips_tables_and_sanitises_names():
    body = {"filename": "lbo summary/2026", "sheets": [
        {"name": "P&L", "columns": ["Line", "Y1", "Y2"], "rows": [["Revenue", 250.0, 262.5], ["EBITDA", 65.5, None]]},
        {"name": "Debt: schedule [senior]", "columns": ["Year", "Balance"], "rows": [[1, 420.0], [2]]},
        {"name": "P&L", "columns": ["x"], "rows": []},
    ]}
    r = client.post("/api/export/workbook", json=body)
    books = _read(r)
    assert 'filename="lbo_summary_2026.xlsx"' in r.headers["content-disposition"]
    assert list(books) == ["P&L", "Debt  schedule  senior", "P&L 2"]
    pl = books["P&L"]
    assert pl.columns.tolist() == ["Line", "Y1", "Y2"]
    assert pl.loc[0, "Y2"] == 262.5 and pd.isna(pl.loc[1, "Y2"])
    assert pd.isna(books["Debt  schedule  senior"].loc[1, "Balance"])


def test_workbook_validation():
    assert client.post("/api/export/workbook", json={"sheets": []}).status_code == 422
    assert client.post("/api/export/workbook", json={"sheets": [{"name": "a", "columns": []}]}).status_code == 422


def test_montecarlo_sample_matches_the_seeded_run():
    req = {"mc": {"n": 4000}, "seed": 11}
    run = client.post("/api/montecarlo/run", json=req).json()
    books = _read(client.post("/api/export/montecarlo-sample", json=req))
    paths = books["Paths"]
    assert len(paths) == 4000 and {"IRR", "MOIC", "Growth", "Exit Multiple"} <= set(paths.columns)
    assert abs(paths["IRR"].mean() - run["summary"]["mean_irr"]) < 1e-9
    about = dict(zip(books["About"]["Item"], books["About"]["Value"]))
    assert about["Seed"] == 11 and about["Paths simulated"] == 4000


# ---------------------------------------------------------------------------
# Per-cell number formats (PLAN.md 2.3a). Excel shows a standard format with
# the reader's own separators (1.234,5 in German Excel), so the workbook holds
# numbers and formats, never text; lakh and crore need patterns of their own.
# ---------------------------------------------------------------------------
def _cells(resp, sheet):
    import openpyxl
    assert resp.status_code == 200, resp.text
    ws = openpyxl.load_workbook(io.BytesIO(resp.content))[sheet]
    return {c.coordinate: (c.value, c.number_format) for row in ws.iter_rows() for c in row}


FORMATTED = {"filename": "f.xlsx", "sheets": [{
    "name": "P&L", "columns": ["", "Y1", "Y2"],
    "rows": [["Revenue", 1234567.25, -2500.5], ["IRR", 0.2116, 0.1], ["MOIC", 2.35, 1.5],
             ["Paths", 50000, 3], ["Note", "text", None]],
    "row_formats": ["money", "percent", "multiple", "integer", "text"],
}]}


def test_each_row_gets_its_number_format_and_stays_a_number():
    cells = _cells(client.post("/api/export/workbook", json=FORMATTED), "P&L")
    assert cells["B2"] == (1234567.25, "#,##0.0")
    assert cells["C2"] == (-2500.5, "#,##0.0")
    assert cells["B3"] == (0.2116, "0.0%")        # a fraction, shown as 21.2%
    assert cells["B4"] == (2.35, '0.00"x"')
    assert cells["B5"] == (50000, "#,##0")
    assert cells["B6"] == ("text", "@")
    # Labels and headers keep Excel's default
    assert cells["A2"][1] == "General" and cells["B1"] == ("Y1", "General")


def test_column_formats_apply_down_a_column_and_a_row_format_wins():
    body = {"sheets": [{"name": "Debt", "columns": ["Year", "Closing", "Rate"],
                        "rows": [[1, 420.0, 0.065], [2, 400.0, 0.07]],
                        "column_formats": ["integer", "money", "percent"]}]}
    cells = _cells(client.post("/api/export/workbook", json=body), "Debt")
    assert cells["A2"][1] == "#,##0" and cells["B3"][1] == "#,##0.0" and cells["C3"] == (0.07, "0.0%")
    body["sheets"][0]["row_formats"] = [None, "text"]
    cells = _cells(client.post("/api/export/workbook", json=body), "Debt")
    assert cells["B2"][1] == "#,##0.0" and cells["B3"][1] == "@"


@pytest.mark.parametrize("value, pattern", [
    (999.5, "#,##0.0"),
    (123456.7, r"#\,##\,##0.0"),              # 1,23,456.7
    (12345678.9, r"#\,##\,##\,##0.0"),        # 1,23,45,678.9
    (-12345678.9, r"#\,##\,##\,##0.0"),       # one section, so Excel keeps the minus sign
    (1234567890123.0, r"##\,##\,##\,##\,##\,##0.0"),   # 12,34,56,78,90,123.0
])
def test_lakh_grouping_writes_a_pattern_sized_to_each_value(value, pattern):
    body = {"grouping": "lakh", "sheets": [{"name": "S", "columns": ["", "Y1"], "rows": [["Revenue", value]],
                                            "row_formats": ["money"]}]}
    cells = _cells(client.post("/api/export/workbook", json=body), "S")
    assert cells["B2"] == (value, pattern)


def _render(pattern: str, value: float) -> str:
    """What Excel shows for one of these single-section patterns: digits fill
    the placeholders from the right, an escaped comma prints as it is."""
    whole, _, frac = pattern.partition(".")
    digits = f"{abs(value):.{len(frac)}f}"
    int_digits, _, frac_digits = digits.partition(".")
    if "\\," not in whole:                      # Excel's own thousands grouping
        shown = f"{int(int_digits):,}"
    else:
        shown, pool = "", list(int_digits)
        for token in reversed(whole.replace("\\,", "|")):
            if token in "#0" and pool:
                shown = pool.pop() + shown
            elif token == "|":
                shown = "," + shown
        shown = "".join(pool) + shown
    return ("-" if value < 0 else "") + shown + ("." + frac_digits if frac else "")


@pytest.mark.parametrize("value, decimals, shown", [
    (12345678.9, 1, "1,23,45,678.9"),
    (-100000.0, 1, "-1,00,000.0"),
    (99999.96, 1, "1,00,000.0"),       # sized after rounding
    (50000, 0, "50,000"),
    (1234567890123.0, 1, "12,34,56,78,90,123.0"),
])
def test_lakh_patterns_show_indian_grouping(value, decimals, shown):
    from api.routers.export import lakh_pattern
    assert _render(lakh_pattern(value, decimals), value) == shown


def test_thousands_grouping_is_the_standard_format():
    body = {"grouping": "thousands", **FORMATTED}
    assert _cells(client.post("/api/export/workbook", json=body), "P&L")["B2"][1] == "#,##0.0"


def test_formats_are_checked():
    bad = {"sheets": [{"name": "S", "columns": ["a"], "rows": [[1]], "row_formats": ["currency"]}]}
    assert client.post("/api/export/workbook", json=bad).status_code == 422
    assert client.post("/api/export/workbook", json={**FORMATTED, "grouping": "crore"}).status_code == 422


def test_the_montecarlo_sample_formats_its_rates_and_multiples():
    import openpyxl
    resp = client.post("/api/export/montecarlo-sample", json={"mc": {"n": 1000}, "seed": 3})
    ws = openpyxl.load_workbook(io.BytesIO(resp.content))["Paths"]
    header = {c.value: c.column_letter for c in ws[1]}
    assert ws[f"{header['IRR']}2"].number_format == "0.0%"
    assert ws[f"{header['MOIC']}2"].number_format == '0.00"x"'
    assert ws[f"{header['Growth']}2"].number_format == "0.0%"
    assert ws[f"{header['Exit Multiple']}2"].number_format == '0.00"x"'
