"""Sector models through the API (PLAN.md R4): the catalogue and its
examples, a run that is the evaluator's answer, a change that never breaks
the model, the workbook and the add-in's cells holding the same cells, and
the model's past filled from accepted quarters (only the caller's).
"""
from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.main import app
from research import compose
from research.definition import Company
from research.evaluate import evaluate
from research.models import CATALOGUE, occupancy
from tests import research_examples
from tests import research_support as support
from tests.addin_cells_check import differences
from tests.test_research_actuals import H1, BALANCE, period, results
from tests.test_uploads import uploads  # noqa: F401 -- the fixture

client = TestClient(app)


def example(n=1):
    return client.get("/api/research/models").json()["examples"][n]


def test_the_examples_shipped_are_the_companies_solved_by_hand():
    assert research_examples.OUT.read_text(encoding="utf-8") == research_examples.text(), \
        "stale: python -m tests.research_examples"


def test_the_catalogue_offers_only_models_solved_by_hand_and_says_none_is_vetted():
    body = client.get("/api/research/models").json()
    models = {m["id"]: m for m in body["models"]}
    assert models["occupancy"]["offered"] is True and models["general"]["offered"] is False
    assert not any(m["vetted"] for m in models.values())
    assert models["occupancy"]["sectors"] == ["hospitals", "hotels"]
    assert {b["id"] for b in body["blocks"]} == {"capacity", "units", "price_volume", "subscribers", "people",
                                                 "order_book", "growth", "share"}
    assert [e["id"] for e in body["examples"]] == ["yatharth-hospital", "lemon-tree-hotels"]


@pytest.mark.parametrize("n, fixture", [(0, "hospitals"), (1, "hotels")])
def test_a_run_answers_the_evaluators_figures_for_every_cell(n, fixture):
    ex = example(n)
    answer = client.post("/api/research/models/run", json={"structure": ex["structure"], "company": ex["company"]})
    assert answer.status_code == 200, answer.text
    body = answer.json()
    values = evaluate(occupancy.MODEL, support.company(support.load(fixture))).values
    compared = 0
    for row in body["rows"]:
        if row["key"] is None:
            continue
        for p, cell in row["cells"].items():
            assert cell["value"] == pytest.approx(values[row["key"]][p]), (row["key"], p)
            compared += 1
    assert compared > 450
    assert body["missing"] == [] and body["dropped"] == [] and body["vetted"] is False
    rev = next(r for r in body["rows"] if r["key"] == "s1_rev")
    assert rev["cells"]["y1"]["formula"] == "s1_avg * days * s1_occ * s1_price / unit_rupees"
    assert rev["cells"]["ly"]["choice"] == "line1_past"
    occ = next(r for r in body["rows"] if r["key"] == "s1_occ")
    assert occ["cells"]["y1"]["kind"] == "assumption" and occ["cells"]["ly"]["kind"] == "figure"


def test_a_run_with_assumptions_still_to_type_works_the_past_out_and_lists_them():
    ex = example()
    company = {**ex["company"], "values": {k: {p: v for p, v in by.items() if p in ("ly", "ytd")}
                                           for k, by in ex["company"]["values"].items()}}
    body = client.post("/api/research/models/run", json={"structure": ex["structure"], "company": company}).json()
    pat = next(r for r in body["rows"] if r["key"] == "pat_own")
    assert pat["cells"]["ly"]["value"] == pytest.approx(19658.11) and pat["cells"]["y1"]["value"] is None
    assert {"row": "s1_occ", "period": "y1"} in body["missing"]


def test_a_change_mends_the_structure_and_a_run_drops_what_no_longer_fits():
    ex = example()
    changed = client.post("/api/research/models/change", json={"structure": ex["structure"], "action": "swap",
                                                                 "key": "s2", "block": "growth"})
    assert changed.status_code == 200
    structure = changed.json()["structure"]
    assert structure["segments"][1]["block"] == "growth"
    body = client.post("/api/research/models/run", json={"structure": structure, "company": ex["company"]}).json()
    assert {"row": "s2_add", "period": "y1"} in body["dropped"]
    assert {"row": "s2_g", "period": "y1"} in body["missing"]
    assert any("not among the companies" in note for note in body["notes"])
    dropped = client.post("/api/research/models/change", json={"structure": ex["structure"], "action": "drop", "key": "s1"})
    assert dropped.json()["structure"]["places"] == "s2"
    added = client.post("/api/research/models/change", json={"structure": ex["structure"], "action": "add", "block": "order_book"})
    assert [s["block"] for s in added.json()["structure"]["segments"]][-1] == "order_book"
    none = client.post("/api/research/models/change", json={"structure": ex["structure"], "action": "places"})
    assert none.json()["structure"]["places"] is None


@pytest.mark.parametrize("structure, code", [
    ({"model": "nope", "segments": [{"key": "s1", "block": "growth"}]}, "model_refused"),
    ({"model": "occupancy", "segments": [{"key": "s1", "block": "share", "of": "s2"}, {"key": "s2", "block": "growth"}]},
     "model_refused"),
])
def test_a_structure_that_would_not_build_is_a_422_with_its_reason(structure, code):
    answer = client.post("/api/research/models/run", json={"structure": structure, "company": example()["company"]})
    assert answer.status_code == 422 and answer.json()["detail"]["code"] == code
    assert answer.json()["detail"]["message"]


def test_a_company_that_does_not_fit_is_a_422_and_bad_shapes_never_reach_the_model():
    ex = example()
    bad = {**ex["company"], "choices": {"line1_past": "guess", "line1_end_to_date": "typed"}}
    answer = client.post("/api/research/models/run", json={"structure": ex["structure"], "company": bad})
    assert answer.status_code == 422 and answer.json()["detail"]["code"] == "company_refused"
    huge = {**ex["company"], "name": "x" * 201}
    assert client.post("/api/research/models/run", json={"structure": ex["structure"], "company": huge}).status_code == 422
    nine = {**ex["structure"], "segments": [{"key": "s%d" % i, "block": "growth"} for i in range(1, 10)]}
    assert client.post("/api/research/models/run", json={"structure": nine, "company": ex["company"]}).status_code == 422


def as_written(answer: dict) -> dict:
    from openpyxl.utils import get_column_letter
    sheets = []
    for sheet in answer["sheets"]:
        cells = {}
        for cell in sheet["cells"]:
            spec = {}
            if cell.get("formula"):
                spec["formula"] = cell["formula"]
            elif cell.get("text") is not None:
                spec["value"] = cell["text"]
            elif cell.get("number") is not None:
                spec["value"] = cell["number"]
            for key, out in (("number_format", "format"), ("bold", "bold"), ("note", "note"), ("fill", "fill")):
                if cell.get(key):
                    spec[out] = cell[key]
            if spec.keys() <= {"format", "bold"}:           # an empty cell with only a format: not a cell to Excel
                continue
            cells[f"{get_column_letter(cell['column'])}{cell['row']}"] = spec
        sheets.append({"name": sheet["title"], "cells": cells, "widths": sheet["widths"], "freeze": sheet["freeze"],
                       "names": sheet["names"]})
    return {"sheets": sheets}


@pytest.mark.parametrize("blank", [False, True])
def test_the_add_in_is_given_the_same_cells_as_the_download(blank):
    ex = example()
    company = ex["company"]
    if blank:                                        # assumptions to type in Excel: blank yellow cells, still named
        company = {**company, "values": {k: {p: v for p, v in by.items() if p in ("ly", "ytd")}
                                         for k, by in company["values"].items()}}
    payload = {"structure": ex["structure"], "company": company}
    cells = client.post("/api/research/models/cells", json=payload)
    download = client.post("/api/research/models/workbook", json=payload)
    assert cells.status_code == 200 and download.status_code == 200, (cells.text, download.text)
    wb = load_workbook(io.BytesIO(download.content))
    body = cells.json()
    problems, count = differences(wb, as_written(body))
    assert problems == []
    assert count > 700
    named = body["sheets"][0]["names"]
    assert len(named) == 100 and named["unit_rupees"] == "$C$4"           # 99 assumption cells and the constant
    yellow = [c for c in body["sheets"][0]["cells"] if c.get("fill") == "FFF8D6"]
    assert len(yellow) == 99 and (all(c.get("number") is None for c in yellow) if blank else True)
    changed = as_written(body)
    first = next(spec for spec in changed["sheets"][0]["cells"].values() if spec.get("fill") == "FFF8D6")
    first["fill"] = "FFFFFF"
    assert differences(wb, changed)[0]                 # and the comparison notices a changed fill


def test_a_blank_workbook_works_out_once_the_analyst_types_the_assumptions_in_it():
    """Load the past only, then type the example's assumptions into the yellow cells: the workbook's own formulas
    give the server's figures, with no call to it."""
    ex = example()
    past = {**ex["company"], "values": {k: {p: v for p, v in by.items() if p in ("ly", "ytd")}
                                        for k, by in ex["company"]["values"].items()}}
    wb = load_workbook(io.BytesIO(client.post("/api/research/models/workbook",
                                              json={"structure": ex["structure"], "company": past}).content))
    names = {n: wb.defined_names[n].attr_text.split("!")[1].replace("$", "") for n in wb.defined_names}
    for row, by in ex["company"]["values"].items():
        for p, v in by.items():
            if "%s_%s" % (row, p) in names:
                wb["Model"][names["%s_%s" % (row, p)]].value = v
    get = support.workbook_values(wb)
    from research.workbook import layout
    where = layout(occupancy.MODEL)
    whole = evaluate(occupancy.MODEL, support.company(support.load("hotels"))).values
    assert get(where.cell("pat_own", "y2")) == pytest.approx(whole["pat_own"]["y2"], abs=0.01)
    assert get(where.cell("chk_bs", "y2")) == pytest.approx(0, abs=1e-6)


# ------------------------------------------------------------------ the past from accepted quarters

def test_the_fill_reads_the_callers_accepted_quarters_and_nobody_elses(uploads):  # noqa: F811
    from db import coverage as rows
    made = client.post("/api/research/companies", json={"name": "Hotel Co"})
    assert made.status_code == 201
    company_id = made.json()["id"]
    stored = period([results(H1, [1000, 900, 950, 1900, 1800, 3800]), BALANCE])
    rows.save_periods("test:alice", __import__("uuid").UUID(company_id),
                      [{"period_end": "2026-09-30", "months": 6, **stored}], None)
    answer = client.get(f"/api/research/companies/{company_id}/fill")
    assert answer.status_code == 200, answer.text
    body = answer.json()
    assert body["values"]["rev"] == {"ytd": 1900, "ly": 3800} and body["unit_rupees"] == 100000
    assert "leases_in_finance_costs" in body["notes"]
    assert client.get(f"/api/research/companies/{company_id}/fill?scope=standalone").json()["notes"] == ["nothing_accepted"]
    uploads("test:bob")
    assert client.get(f"/api/research/companies/{company_id}/fill").status_code == 404
