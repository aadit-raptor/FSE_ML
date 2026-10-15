"""Check a downloaded checked workbook against the stored period it came from
(PLAN.md R2), for the browser tests: ``python -m tests.coverage_workbook_check
WORKBOOK PERIOD_JSON``.

Every figure the period holds is worked out from the workbook's own formulas
(``tests/research_support.py``'s reader of Excel formulas, which shares no
code with the writer) and must equal the stored figure. Prints how many
figures and formulas it checked; exits 1 on the first difference.
"""
from __future__ import annotations

import json
import sys

from openpyxl import load_workbook

from research.coverage import FIRST_COLUMN, FIRST_ROW
from tests.research_support import workbook_values


def check(workbook_path: str, period_path: str) -> dict:
    period = json.loads(open(period_path, encoding="utf-8").read())
    wb = load_workbook(workbook_path)
    sheets = [s for s in wb.sheetnames if s != "About"]
    figures = formulas = 0
    for sheet, table in zip(sheets, period["figures"]["tables"], strict=True):
        worked = workbook_values(wb, sheet)
        for r, row in enumerate(table["rows"]):
            for c, value in enumerate(row["values"]):
                if value is None:
                    continue
                cell = wb[sheet].cell(FIRST_ROW + r, FIRST_COLUMN + c)
                got = worked(cell.coordinate)
                if abs(got - value) > 1e-6:
                    raise SystemExit(f"{sheet}!{cell.coordinate}: the workbook gives {got}, the period holds {value}")
                figures += 1
                formulas += isinstance(cell.value, str) and cell.value.startswith("=")
    return {"figures": figures, "formulas": formulas}


if __name__ == "__main__":
    print(json.dumps(check(sys.argv[1], sys.argv[2])))
