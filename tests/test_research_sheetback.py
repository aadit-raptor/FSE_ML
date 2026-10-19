"""Saving a model back from Excel (PLAN.md R5): the named cells read back,
and a sheet whose structure was changed refused rather than guessed.

The sheet is the workbook Variater writes (research/workbook.py `write`),
opened with openpyxl and read as the add-in reads Excel: values from A1 and
the sheet's defined names with their addresses.
"""
from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from research.sheetback import NOT_A_NUMBER, OTHER_MODEL, RESTRUCTURED, SheetRefused, read_back
from research.workbook import SAVED_NAME, write
from tests.test_research_rollforward import yatharth

SAVED = "5b0d1f3e-6a8e-4d55-9b1c-2f7f0e9a1c11"


def workbook(model, company, saved=SAVED):
    buf = io.BytesIO()
    write(model, company, allow_missing=True, saved=saved).save(buf)
    return load_workbook(io.BytesIO(buf.getvalue()))


def as_excel_reads_it(wb):
    """The Model sheet's values from A1 (formulas stand in as their text: only words and named cells are read)
    and its names, as {name: address}."""
    ws = wb["Model"]
    grid = [[cell.value if cell.value is not None else "" for cell in row]
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]
    names = {}
    for name, defined in wb.defined_names.items():
        sheet, address = next(defined.destinations)
        if sheet == "Model":
            names[name] = address
    return grid, names


def test_a_sheet_saved_unchanged_gives_the_same_company():
    model, company = yatharth()
    grid, names = as_excel_reads_it(workbook(model, company))
    back = read_back(model, company, SAVED, grid, names)
    assert back.changed == [] and back.company.values == company.values
    assert back.company.constants == company.constants


def test_assumptions_and_constants_typed_in_excel_come_back():
    model, company = yatharth()
    wb = workbook(model, company)
    ws = wb["Model"]
    names = {n: next(d.destinations)[1] for n, d in wb.defined_names.items()}
    ws[names["mat_sh_y1"].replace("$", "")] = 0.31
    ws[names["s1_add_roy"].replace("$", "")] = None              # cleared: an assumption not typed
    ws[names["unit_rupees"].replace("$", "")] = 10_000_000
    grid, found = as_excel_reads_it(wb)
    back = read_back(model, company, SAVED, grid, found)
    assert back.company.values["mat_sh"]["y1"] == 0.31
    assert "roy" not in back.company.values.get("s1_add", {})
    assert back.company.constants["unit_rupees"] == 10_000_000
    assert ("mat_sh", "y1", company.values["mat_sh"]["y1"], 0.31) in back.changed
    # figures typed from the documents are not read back: they stay as Variater holds them
    assert back.company.values["rev"] == company.values["rev"]


def _insert_row(grid, names, at):
    grid = grid[: at - 1] + [[""] * len(grid[0])] + grid[at - 1:]
    moved = {}
    for name, address in names.items():
        letters = "".join(ch for ch in address if ch.isalpha())
        row = int("".join(ch for ch in address if ch.isdigit()))
        moved[name] = "$%s$%d" % (letters, row + 1 if row >= at else row)     # as Excel moves a name
    return grid, moved


def test_a_row_inserted_in_the_sheet_is_refused():
    model, company = yatharth()
    grid, names = as_excel_reads_it(workbook(model, company))
    grid, names = _insert_row(grid, names, 20)
    with pytest.raises(SheetRefused) as refused:
        read_back(model, company, SAVED, grid, names)
    assert refused.value.code == RESTRUCTURED and "row 20" in str(refused.value)


def test_a_named_cell_moved_or_removed_is_refused():
    model, company = yatharth()
    grid, names = as_excel_reads_it(workbook(model, company))
    moved = dict(names, mat_sh_y1="$" + get_column_letter(12) + "$500")
    with pytest.raises(SheetRefused) as refused:
        read_back(model, company, SAVED, grid, moved)
    assert refused.value.code == RESTRUCTURED and "mat_sh_y1 has moved" in str(refused.value)
    gone = {k: v for k, v in names.items() if k != "tax_rate_y2"}
    with pytest.raises(SheetRefused) as refused:
        read_back(model, company, SAVED, grid, gone)
    assert "tax_rate_y2 is gone" in str(refused.value)


def test_words_in_an_assumption_are_refused():
    model, company = yatharth()
    wb = workbook(model, company)
    names = {n: next(d.destinations)[1] for n, d in wb.defined_names.items()}
    wb["Model"][names["mat_sh_y1"].replace("$", "")] = "about a third"
    grid, found = as_excel_reads_it(wb)
    with pytest.raises(SheetRefused) as refused:
        read_back(model, company, SAVED, grid, found)
    assert refused.value.code == NOT_A_NUMBER and "about a third" in str(refused.value)


def test_another_models_sheet_is_refused():
    model, company = yatharth()
    grid, names = as_excel_reads_it(workbook(model, company, saved="0" * 8 + SAVED[8:]))
    with pytest.raises(SheetRefused) as refused:
        read_back(model, company, SAVED, grid, names)
    assert refused.value.code == OTHER_MODEL
    grid, names = as_excel_reads_it(workbook(model, company))
    with pytest.raises(SheetRefused):
        read_back(model, company, SAVED, grid, {k: v for k, v in names.items() if k != SAVED_NAME})


def test_the_saved_models_reference_is_beside_the_title_and_named():
    model, company = yatharth()
    wb = workbook(model, company)
    assert wb["Model"]["B1"].value == "Saved in Variater: " + SAVED
    assert next(wb.defined_names[SAVED_NAME].destinations) == ("Model", "$B$1")
    unsaved = workbook(model, company, saved=None)
    assert unsaved["Model"]["B1"].value is None and SAVED_NAME not in unsaved.defined_names
