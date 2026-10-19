"""Whether what the Excel add-in wrote holds the same cells as the checked
workbook's download (PLAN.md R3, "Done when").

``written`` is what Excel was asked to hold, in the add-in's own terms (the
browser tests record it from the pane, ``web/e2e/excel.spec.ts``)::

    {"sheets": [{"name": "Results consolidated",
                 "cells": {"C6": {"formula": "=C7+C8"} | {"value": 12.5 | "words"},
                           ... each also "format", "bold", "note", "fill"},
                 "names": {"s1_occ_y1": "$G$20"},
                 "widths": {"A": 48}, "freeze": "C6"}]}

Every cell of the download must be written with the same value or formula,
number format, boldness, note and fill (a blank cell counts when it is
filled: a sector model's assumption not typed yet), every defined name on the
sheet must be written at the same address, every written cell must be in the
download, and the sheets must come in the same order. The one cell allowed
to differ is when each was generated (About, row 3). Run as::

    python -m tests.addin_cells_check workbook.xlsx written.json

which prints ``{"cells": n, "sheets": n, "differences": [...]}`` and exits 1
when there is a difference.
"""
from __future__ import annotations

import json
import sys
from typing import Any, Mapping

from openpyxl import load_workbook

from tests import research_support as support

# When each was made: the download and the add-in's write are two requests
GENERATED = ("About", "B3")


def _fill(cell) -> str | None:
    if cell.fill is not None and cell.fill.patternType == "solid":
        return str(cell.fill.fgColor.rgb)[-6:]
    return None


def _expected(ws) -> dict:
    out = {}
    for row in ws.iter_rows():
        for cell in row:
            fill = _fill(cell)
            if cell.value is None and fill is None:
                continue
            spec: dict[str, Any] = {}
            if cell.data_type == "f":
                spec["formula"] = cell.value
            elif cell.value is not None:
                spec["value"] = support.as_excel_holds_it(cell.value)
            if fill:
                spec["fill"] = fill
            if cell.number_format and cell.number_format != "General":
                spec["format"] = cell.number_format
            if cell.font is not None and cell.font.b:
                spec["bold"] = True
            if cell.comment is not None:
                spec["note"] = cell.comment.text
            out[cell.coordinate] = spec
    return out


def _same(a: Any, b: Any) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(float(a) - float(b)) <= 1e-9 * max(1.0, abs(float(a)))
    return a == b


def differences(wb, written: Mapping) -> tuple[list[str], int]:
    problems: list[str] = []
    sheets = list(written.get("sheets") or [])
    names = [s["name"] for s in sheets]
    if names != wb.sheetnames:
        problems.append(f"sheets {names} != {wb.sheetnames}")
    count = 0
    for sheet in sheets:
        if sheet["name"] not in wb.sheetnames:
            continue
        ws = wb[sheet["name"]]
        expected = _expected(ws)
        # a cell given only a format or a font holds nothing to Excel, as an empty cell of the download holds nothing
        got = {a: c for a, c in (sheet.get("cells") or {}).items()
               if any(c.get(k) is not None for k in ("formula", "value", "note", "fill"))}
        for address in sorted(set(expected) | set(got)):
            want, have = expected.get(address), got.get(address)
            if want is None or have is None:
                problems.append(f"{sheet['name']}!{address}: {'not in the download' if want is None else 'not written'}")
                continue
            count += 1
            if isinstance(have.get("value"), str) and have.get("format") == "@":
                # words are written into a cell formatted as text, so one that starts with = stays words
                have = {k: v for k, v in have.items() if k != "format"}
            for key in ("formula", "value", "format", "bold", "note", "fill"):
                if (sheet["name"], address) == GENERATED and key == "value":
                    continue
                if not _same(want.get(key), have.get(key)):
                    problems.append(f"{sheet['name']}!{address} {key}: {have.get(key)!r} != {want.get(key)!r}")
        for letter, width in (sheet.get("widths") or {}).items():
            if not _same(ws.column_dimensions[letter].width, width):
                problems.append(f"{sheet['name']} column {letter} width {width} != {ws.column_dimensions[letter].width}")
        if (sheet.get("freeze") or None) != ws.freeze_panes:
            problems.append(f"{sheet['name']} freeze {sheet.get('freeze')!r} != {ws.freeze_panes!r}")
        named = {name: wb.defined_names[name].attr_text.split("!", 1)[1] for name in wb.defined_names
                 if wb.defined_names[name].attr_text.split("!", 1)[0].strip("'") == sheet["name"]}
        if named != (sheet.get("names") or {}):
            problems.append(f"{sheet['name']} names: {len(sheet.get('names') or {})} written, {len(named)} in the download")
    return problems, count


def main(argv: list[str]) -> int:
    workbook_path, written_path = argv
    wb = load_workbook(workbook_path)
    with open(written_path, encoding="utf-8") as fh:
        written = json.load(fh)
    problems, count = differences(wb, written)
    print(json.dumps({"cells": count, "sheets": len(written.get("sheets") or []), "differences": problems[:50]}))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
